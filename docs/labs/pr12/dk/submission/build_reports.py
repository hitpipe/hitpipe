"""Build the four illustrated reports; preserve existing title-page contents."""
from pathlib import Path
from datetime import datetime,timezone
import re,json,hashlib,copy
from lxml import etree
from zipfile import ZipFile
from docx import Document
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent
BASE=HERE.parent
FIGURES=json.loads((HERE/'figures.json').read_text())
CAPTURES=json.loads((HERE/'capture-manifest.json').read_text())
BODY_START='http://schemas.openxmlformats.org/wordprocessingml/2006/main'

def font(run):
    run.font.name='Times New Roman';run.font.size=Pt(14)
    fonts=run._r.get_or_add_rPr().find(qn('w:rFonts'))
    for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Times New Roman')
    lang=run._r.get_or_add_rPr().find(qn('w:lang'))
    if lang is None:lang=OxmlElement('w:lang');run._r.get_or_add_rPr().append(lang)
    lang.set(qn('w:val'),'ru-RU')

def paragraph_format(p,kind='body'):
    f=p.paragraph_format;f.line_spacing=1.5;f.space_before=Pt(0);f.space_after=Pt(0)
    f.left_indent=Pt(0);f.right_indent=Pt(0);f.widow_control=True
    f.first_line_indent=Cm(1.25) if kind=='body' else Pt(0)
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if kind in ['image','caption']:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if kind=='heading':f.space_before=Pt(12);f.space_after=Pt(6);f.keep_with_next=True;f.keep_together=True
    if kind=='caption':f.space_before=Pt(6);f.space_after=Pt(12);f.keep_together=True;f.keep_with_next=False
    for r in p.runs:font(r)

def inline(p,text):
    pat=r'(\[[^\]]+\]\(https?://[^)]+\)|`[^`]+`|\*\*[^*]+\*\*)'
    for part in re.split(pat,text):
        match=re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)',part)
        if match:
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(match[2],'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
            r=OxmlElement('w:r');props=OxmlElement('w:rPr')
            fonts=OxmlElement('w:rFonts')
            for k in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+k),'Times New Roman')
            props.append(fonts);size=OxmlElement('w:sz');size.set(qn('w:val'),'28');props.append(size)
            color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');props.append(color);r.append(props)
            t=OxmlElement('w:t');t.text=match[1];r.append(t);h.append(r);p._p.append(h)
        else:
            value=part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('`') else part
            run=p.add_run(value);font(run)
            if part.startswith('**'):run.bold=True

def p(doc,text,kind='body',style=None):
    para=doc.add_paragraph(style=style);inline(para,text);paragraph_format(para,kind);return para

def borders(table):
    prop=table._tbl.tblPr
    b=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        node=OxmlElement('w:'+edge)
        for key,val in [('val','single'),('sz','8'),('color','000000')]:node.set(qn('w:'+key),val)
        b.append(node)
    prop.append(b)

def table(doc,rows,record_header=False):
    sec=doc.sections[-1];width=sec.page_width-sec.left_margin-sec.right_margin
    t=doc.add_table(rows=0,cols=2);t.autofit=False
    ratios=[0.30,0.70];widths=[int(width*x)for x in ratios]
    for c,w in zip(t.columns,widths):c.width=w
    borders(t)
    for i,values in enumerate(rows):
        row=t.add_row();pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for cell,w,text in zip(row.cells,widths,values):
            cell.width=w;inline(cell.paragraphs[0],text);paragraph_format(cell.paragraphs[0],'cell')
            if i==0:
                for run in cell.paragraphs[0].runs:run.bold=True
    return t

def parse_table_line(line):
    result=[];buf='';code=False;escaped=False
    for ch in line.strip().strip('|'):
        if escaped:buf+=ch;escaped=False;continue
        if ch=='\\':escaped=True;buf+=ch;continue
        if ch=='`':code=not code
        if ch=='|' and not code:result.append(buf.strip());buf=''
        else:buf+=ch
    result.append(buf.strip());return result

def render_table(doc,rows):
    if len(rows[0])==2:table(doc,rows);return
    # Wide source tables become labelled records at 14 pt, without reducing type
    # size or losing any original field. Every record repeats its own header.
    headers=rows[0]
    for record in rows[1:]:
        if len(record)!=len(headers):raise ValueError(('Table cell count',headers,record))
        pairs=[[headers[0],record[0]]]+[[label,value]for label,value in zip(headers[1:],record[1:])]
        table(doc,pairs,True)
        p(doc,'','cell')

def figure(doc,identity):
    f=FIGURES[identity];capture=CAPTURES[identity]
    caption=re.sub(r'^(Рисунок\s+\d+)\s*[—–-]\s*',r'\1 - ',f['caption'])
    if not capture.get('images'):
        raise ValueError(f'{identity}: no actual capture; do not insert a missing-result placeholder into the report')
    for i,item in enumerate(capture['images']):
        image=HERE/'evidence'/item['file']
        assert hashlib.sha256(image.read_bytes()).hexdigest()==item['sha256']
        para=p(doc,'','image');para.paragraph_format.keep_with_next=True
        run=para.add_run();font(run)
        run.add_picture(str(image),width=Cm(16))
        suffix=f' (часть {i+1} из {len(capture["images"])})' if len(capture['images'])>1 else ''
        p(doc,caption+suffix,'caption',style='Caption')

def markdown(doc,text):
    lines=text.splitlines();i=0;code=False
    while i<len(lines):
        raw=lines[i];line=raw.strip();i+=1
        if line.startswith('```'):code=not code;continue
        if not line:continue
        if code:p(doc,raw,'cell');continue
        marker=re.fullmatch(r'<!-- FIGURE:([a-z0-9-]+) -->',line)
        if marker:figure(doc,marker[1]);continue
        h=re.match(r'^(#{1,4}) (.+)$',line)
        if h:
            level=min(len(h[1])-1,3)
            if level:
                # A spacer after the final table must not consume a blank page
                # immediately before a chapter with pageBreakBefore.
                body=doc._element.body
                while len(body)>1:
                    node=list(body)[-2] if list(body)[-1].tag==qn('w:sectPr') else list(body)[-1]
                    if node.tag!=qn('w:p') or node.xpath('.//w:t | .//w:drawing | .//w:br'):break
                    body.remove(node)
                heading=doc.add_heading(h[2],level=level);paragraph_format(heading,'heading')
                for r in heading.runs:r.italic=True;r.bold=False
            continue
        if line.startswith('|'):
            rows=[]
            while True:
                cells=parse_table_line(line)
                if not all(re.fullmatch(r':?-+:?',c)for c in cells):rows.append(cells)
                if i==len(lines)or not lines[i].lstrip().startswith('|'):break
                line=lines[i].strip();i+=1
            render_table(doc,rows);continue
        p(doc,('• '+line[2:])if line.startswith('- ')else line)

def build(name):
    original=BASE/'reports'/(name+'-dk.docx')
    doc=Document(original);body=doc._element.body
    cut=None
    for i,node in enumerate(body):
        s=''.join(n.text or ''for n in node.iter(qn('w:t')))
        if s.strip() in ['СОДЕРЖАНИЕ','ВВЕДЕНИЕ']:cut=i;break
    assert cut is not None,name
    cover=[copy.deepcopy(n)for n in list(body)[:cut]]
    cover_hash=hashlib.sha256(b''.join(etree.tostring(n,method='c14n',exclusive=True)for n in cover)).hexdigest()
    for node in list(body)[cut:]:
        if node.tag!=qn('w:sectPr'):body.remove(node)
    # Existing title-page nodes, including unknown student/teacher fields, stay
    # byte-for-byte unchanged. No names, dates, group or signatures are filled.
    for style in doc.styles:
        if style.type in [1,2]:
            style.font.name='Times New Roman';style.font.size=Pt(14)
    normal=doc.styles['Normal'];normal.paragraph_format.line_spacing=1.5;normal.paragraph_format.first_line_indent=Cm(1.25)
    for name_ in ['Heading 1','Heading 2','Heading 3']:
        doc.styles[name_].font.italic=True;doc.styles[name_].font.bold=False
    if name.startswith('0302'):
        title=p(doc,'СОДЕРЖАНИЕ','heading')
        for r in title.runs:r.italic=True
        seed=p(doc,'[ОГЛАВЛЕНИЕ: ОБНОВИТЬ ЧЕРЕЗ render_reports.py]','cell')
        doc.add_page_break()
    markdown(doc,(HERE/(name+'-dk.md')).read_text())
    settings=doc.settings.element;u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');settings.append(u)
    doc.core_properties.created=datetime.now(timezone.utc)
    doc.core_properties.modified=datetime.now(timezone.utc)
    doc.core_properties.author='Дмитрий Карачебан'
    doc.core_properties.last_modified_by='Дмитрий Карачебан'
    doc.core_properties.title={'0104-pz01':'Инструментальный контур','0104-pz02':'Модульная архитектура','0302-pz01':'Требования и регламент','0302-pz02':'Система проверок'}[name]
    doc.core_properties.comments=''
    doc.core_properties.keywords='HitPipe; IAM; Дмитрий Карачебан'
    doc.core_properties.subject='Индивидуальный отчёт о практической работе'
    path=HERE/(name+'-dk.docx');doc.save(path)
    # Check preservation of the actual inherited cover nodes, not just strings.
    after=Document(path)
    after_nodes=list(after._element.body)[:cut]
    after_hash=hashlib.sha256(b''.join(etree.tostring(n,method='c14n',exclusive=True)for n in after_nodes)).hexdigest()
    assert after_hash==cover_hash,(name,'title page changed')
    return {'file':path.name,'sha256_before_renderer':hashlib.sha256(path.read_bytes()).hexdigest(),'cover_nodes':cut,'cover_unchanged_before_renderer':True,'source_title_page':str(original)}

def validate(records):
    from lxml import etree
    result={'reports':[],'image_policy':'Actual browser captures authorized by user 2026-10-19; no generated screenshots','submission_ready':False,'external_opposition':'NOT_CONDUCTED','personal_academic_attestation':False}
    for record in records:
        path=HERE/record['file'];name=path.name.removesuffix('-dk.docx')
        expected=[]
        for identity,capture in CAPTURES.items():
            if identity.startswith(name):expected.extend(capture.get('images',[]))
        with ZipFile(path)as z:
            assert z.testzip()is None
            for member in z.namelist():
                if member.endswith(('.xml','.rels')):etree.fromstring(z.read(member))
            tree=etree.fromstring(z.read('word/document.xml'))
            ns={'w':BODY_START,'pic':'http://schemas.openxmlformats.org/drawingml/2006/picture'}
            text='\n'.join(tree.xpath('//w:t/text()',namespaces=ns))
            assert '[МЕСТО ДЛЯ СКРИНШОТА'not in text
            assert len(tree.xpath('//pic:pic',namespaces=ns))==len(expected),(name,'image count')
            actual_images={hashlib.sha256(z.read(n)).hexdigest()for n in z.namelist()if n.startswith('word/media/')}
            assert {i['sha256']for i in expected}<=actual_images,(name,'embedded original captures changed')
            body=tree.find('w:body',ns);nodes=list(body);start=next(i for i,n in enumerate(nodes)if ''.join(n.xpath('.//w:t/text()',namespaces=ns)).strip()=='ВВЕДЕНИЕ');tables=[t for n in nodes[start:]for t in ([n]if n.tag==qn('w:tbl')else n.xpath('.//w:tbl',namespaces=ns))]
            for t in tables:
                assert t.xpath('./w:tblPr/w:tblBorders',namespaces=ns) or all(c.xpath('./w:tcPr/w:tcBorders/w:top[@w:val="single"]',namespaces=ns) and c.xpath('./w:tcPr/w:tcBorders/w:bottom[@w:val="single"]',namespaces=ns) for c in t.xpath('./w:tr/w:tc',namespaces=ns)),(name,'table borders missing')
            result['reports'].append({'file':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'embedded_captures':len(expected),'tables':len(tables),'zip_xml_images_verified':True})
    return result

if __name__=='__main__':
    records=[build(name)for name in ['0104-pz01','0104-pz02','0302-pz01','0302-pz02']]
    (HERE/'build-validation.json').write_text(json.dumps(records,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(records,ensure_ascii=False,indent=2))

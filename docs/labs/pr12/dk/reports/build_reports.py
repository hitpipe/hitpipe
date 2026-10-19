#!/usr/bin/env python3
"""Build DK's four reports from edited Markdown using the pinned DOCX forms.

Requires python-docx==1.2.0. No screenshot capture, image generation, renderer or
external application is invoked. Human evidence fields remain explicitly pending.
"""
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile
import xml.etree.ElementTree as ET

import docx
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[4]
TEMPLATES=ROOT/'docs/labs/pr12/templates'
FIGURES=json.loads((HERE/'figures.json').read_text())

def inline(paragraph,text):
    pattern=r'(\[[^\]]+\]\(https?://[^)]+\)|`[^`]+`|\*\*[^*]+\*\*)'
    for part in re.split(pattern,text):
        link=re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)',part)
        if link:
            item=OxmlElement('w:hyperlink')
            rid=paragraph.part.relate_to(link[2],'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
            item.set(qn('r:id'),rid)
            run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');props.append(color);run.append(props)
            t=OxmlElement('w:t');t.text=link[1];run.append(t);item.append(run);paragraph._p.append(item)
        else:
            run=paragraph.add_run(part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('`') else part)
            if part.startswith('**'):run.bold=True
            if part.startswith('`'):run.font.name='Consolas';run.font.size=Pt(10)

def paragraph(document,text,style=None):
    p=document.add_paragraph(style=style);inline(p,text);return p

def write_figure_instructions():
    """Keep the human guide and the DOCX slots on the same source of truth."""
    lines=['# DK — 21 человеческий снимок, агент изображения не создаёт\n',
           'Ниже точные версии существующих артефактов. Предварительный код-кандидат не является итоговым submission tag. После human merge/source rewrite использовать новый фактический submission-record и заново снять затронутые результаты; не переносить старый SHA молча. Pending human gate не превращается в PASS наличием placeholder.\n']
    fields=[('insert_location','Место'),('state','Статус'),('navigation','Навигация'),
            ('source_sha','Версия'),('cwd','Каталог'),('command','Действия'),
            ('must_be_visible','В кадре'),('target_file','Файл'),('redaction','Редакция')]
    for identity,figure in FIGURES.items():
        lines.append('## '+identity+' — '+figure['caption']+'\n')
        for key,label in fields:
            value=figure[key]
            lines.append(label+': '+('; '.join(value) if isinstance(value,list) else value)+'\n')
    (HERE.parent/'evidence/screenshots/INSTRUCTIONS.md').write_text('\n'.join(lines))

def format_table(document,table,rows):
    """Keep complete records together and repeat context on every table page.

    Widths use the current printable area; no font shrinking or text removal.
    Teacher cover tables are retained verbatim and do not pass through here.
    """
    section=document.sections[-1]
    available=section.page_width-section.left_margin-section.right_margin
    count=len(table.columns)
    if count==2:
        ratios=[0.28,0.72]
    elif count==6:
        ratios=[0.125,0.255,0.17,0.10,0.11,0.24]
    elif rows[0][0]=='№':
        ratios=[0.07,0.45,0.48]
    elif rows[0][0]=='Mutant / правило':
        ratios=[0.28,0.07,0.32,0.33]
    else:
        ratios=[1/count]*count
    widths=[int(available*ratio) for ratio in ratios]
    table.autofit=False
    for column,width in zip(table.columns,widths):column.width=width
    for n,row in enumerate(table.rows):
        properties=row._tr.get_or_add_trPr()
        properties.append(OxmlElement('w:cantSplit'))
        if n==0:properties.append(OxmlElement('w:tblHeader'))
        for cell,width in zip(row.cells,widths):
            cell.width=width
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.first_line_indent=Pt(0)
                p.paragraph_format.left_indent=Pt(0)
                p.paragraph_format.right_indent=Pt(0)
                p.paragraph_format.space_before=Pt(0)
                p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.line_spacing=1.1

def render_figure(document,identity):
    f=FIGURES[identity]
    document.add_paragraph(f['caption'],style='Caption')
    p=paragraph(document,'[МЕСТО ДЛЯ СКРИНШОТА: '+identity+']')
    for run in p.runs:run.bold=True
    for key,label in [('insert_location','Куда вставить'),('state','Статус'),('navigation','Навигация'),('source_sha','Точная исходная версия'),('cwd','Рабочий каталог'),('command','Команды / действия'),('must_be_visible','Что должно быть видно'),('target_file','Имя файла'),('redaction','Что скрыть')]:
        value=f[key]
        if isinstance(value,list):value='; '.join(value)
        paragraph(document,label+': '+value)

def render_markdown(document,text):
    lines=text.splitlines();i=0;code=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if line.startswith('```'):code=not code;continue
        if code:
            p=document.add_paragraph();run=p.add_run(lines[i-1]);run.font.name='Consolas';run.font.size=Pt(9);continue
        if not line:continue
        marker=re.fullmatch(r'<!-- FIGURE:([a-z0-9-]+) -->',line)
        if marker:render_figure(document,marker[1]);continue
        heading=re.match(r'^(#{1,4}) (.*)',line)
        if heading:
            if len(heading[1])==1:continue  # the preserved teacher title page owns the title
            document.add_heading(heading[2],level=min(3,len(heading[1])-1));continue
        if line.startswith('|'):
            rows=[]
            while True:
                cells=[c.strip() for c in line.strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',c) for c in cells):rows.append(cells)
                if i>=len(lines) or not lines[i].strip().startswith('|'):break
                line=lines[i].strip();i+=1
            table=document.add_table(rows=0,cols=max(map(len,rows)))
            try:table.style='Table Grid'
            except KeyError:pass
            for n,cells in enumerate(rows):
                row=table.add_row()
                for cell,value in zip(row.cells,cells):
                    inline(cell.paragraphs[0],value)
                    for p in cell.paragraphs:
                        for r in p.runs:r.font.size=Pt(10);r.bold=n==0
            format_table(document,table,rows)
            continue
        if line.startswith('- '):paragraph(document,line[2:],style='List Bullet');continue
        paragraph(document,line)

def build(name):
    template=TEMPLATES/(name+'-report-template.docx');document=Document(template)
    body=document._element.body
    cut=None
    for i,node in enumerate(body):
        text=''.join(node.itertext())
        # Inspect actual w:t nodes rather than matching style/TOC text fragments.
        text=''.join(t.text or '' for t in node.iter(qn('w:t')))
        if text in ('СОДЕРЖАНИЕ','ВВЕДЕНИЕ'):cut=i;break
    if cut is None:raise ValueError('Template title/body boundary missing: '+name)
    for node in list(body)[cut:]:
        if node.tag!=qn('w:sectPr'):body.remove(node)
    # Teacher forms split the group label across differently formatted runs.
    for p in body.iter(qn('w:p')):
        nodes=list(p.iter(qn('w:t')))
        label=''.join(n.text or '' for n in nodes)
        if label.strip().startswith('Студент гр. №'):
            nodes[0].text='Студент гр. № [ВПИСАТЬ ГРУППУ]'
            for node in nodes[1:]:node.text=''
    texts=[n for n in body.iter(qn('w:t'))]
    for n in texts:
        if n.text and n.text.strip()=='КАФЕДРА №':n.text='КАФЕДРА № [ВПИСАТЬ КАФЕДРУ]'
    names=[n for n in texts if n.text and 'инициалы, фамилия' in n.text]
    if names:names[-1].text='Д. Карачебан'
    # Remove the old form's trailing blank/page-break paragraphs before making
    # one explicit title/body break, leaving the original title layout intact.
    for node in reversed(list(body)):
        if node.tag==qn('w:sectPr'):continue
        if node.tag==qn('w:p') and not ''.join(n.text or '' for n in node.iter(qn('w:t'))):body.remove(node)
        else:break
    document.add_paragraph('2026').alignment=WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()
    if name.startswith('0302'):
        document.add_paragraph('СОДЕРЖАНИЕ',style='Title')
        p=document.add_paragraph();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-3" \\h \\z \\u')
        run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='[ОГЛАВЛЕНИЕ: ОБНОВИТЬ ЧЕРЕЗ render_reports.py]';run.append(text);field.append(run);p._p.append(field)
        paragraph(document,'При пересборке обновить оглавление через render_reports.py или Word/LibreOffice. После вставки реальных снимков обновить его снова и повторно проверить страницы.')
        document.add_page_break()
    render_markdown(document,(HERE/(name+'-dk.md')).read_text())
    settings=document.settings.element
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    document.core_properties.author='Дмитрий Карачебан (подготовка агентом Codex)'
    document.core_properties.subject='MANUAL_LAB_PR12; PREPARED_FOR_REVIEW; human gates pending'
    document.core_properties.title=name+' — IAM SessionPolicy'
    path=HERE/(name+'-dk.docx');document.save(path)
    return {'file':path.name,'template':template.name,'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()}

def validate(records):
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    seen=[]
    for item in records:
        path=HERE/item['file']
        with ZipFile(path) as archive:
            if archive.testzip() is not None:raise ValueError('Broken ZIP')
            for member in archive.namelist():
                if member.endswith(('.xml','.rels')):ET.fromstring(archive.read(member))
            if any(n.startswith('word/media/') for n in archive.namelist()):raise ValueError('Unexpected image')
            tree=ET.fromstring(archive.read('word/document.xml'))
            text='\n'.join(n.text or '' for n in tree.findall('.//w:t',ns))
            if '{{' in text or 'TEMPLATE_UNFILLED' in text:raise ValueError('Unresolved template content')
            ids=re.findall(r'\[МЕСТО ДЛЯ СКРИНШОТА: ([a-z0-9-]+)\]',text);seen+=ids
            item['screenshot_slots']=len(ids);item['tables']=len(tree.findall('.//w:tbl',ns));item['xml_valid']=True
            for identity in ids:
                f=FIGURES[identity]
                if f['target_file'] not in text or f['source_sha'] not in text:raise ValueError('Missing resolved figure metadata')
        # A second independent parser roundtrip catches missing OOXML relationships.
        Document(path)
    if len(seen)!=21 or len(set(seen))!=21 or set(seen)!=set(FIGURES):raise ValueError('Expected exactly 21 unique figure slots')
    return {'status':'PASS','checks':['ZIP integrity','all XML/rels parse','python-docx reopen','21 unique resolved image placeholders','no captured/generated media','no unresolved template tokens'],
            'python_docx_version':docx.__version__,'reports':records,
            'rendering':'NOT_RUN_BY_BUILDER: run render_reports.py with a real LibreOffice/UNO renderer; see visual-validation.json for the checked DOCX hashes. Revalidate pagination after inserting real screenshots.',
            'submission_ready':False}

if __name__=='__main__':
    write_figure_instructions()
    records=[build(name) for name in ['0302-pz01','0302-pz02','0104-pz01','0104-pz02']]
    result=validate(records);(HERE/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(result,ensure_ascii=False,indent=2))

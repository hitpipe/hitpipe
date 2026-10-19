import sys,json,concurrent.futures
from pathlib import Path
sys.path.insert(0,'/tmp');from hitpipe_dk_review_runner import *
from docx import Document
PY='/Users/tima/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3';RENDER='/Users/tima/.codex/plugins/cache/openai-primary-runtime/documents/26.1015.12040/skills/documents/render_docx.py';audit={}
for f in sorted((C/'docs/labs/pr12/dk/reports').glob('*.docx')):
 d=Document(f);text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells);(E/(f.stem+'-full-text.txt')).write_text(text)
 audit[f.stem]={'sha256':sha(f),'paragraphs':len(d.paragraphs),'tables':len(d.tables),'table_rows':[len(t.rows) for t in d.tables],'cells':sum(len(row.cells) for t in d.tables for row in t.rows),'sections':len(d.sections),'inline_shapes':len(d.inline_shapes)}
(E/'docx-structure.json').write_text(json.dumps(audit,indent=2)+'\n')
def render(f):
 run('render-'+f.stem,[PY,RENDER,str(f),'--output_dir',str(E/'render'/f.stem),'--dpi','100','--emit_pdf'],timeout=300)
with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:list(pool.map(render,sorted((C/'docs/labs/pr12/dk/reports').glob('*.docx'))))
print('Render pages',{p.name:len(list(p.glob('page-*.png'))) for p in (E/'render').iterdir()})

from pathlib import Path
import sys,json,re,zipfile,datetime
from collections import Counter
from lxml import etree
from docx import Document
import pdfplumber
sys.path.insert(0,'/tmp');from hitpipe_dk_review_runner import *
norm=lambda s:''.join(re.findall(r'[A-Za-zА-Яа-яЁё0-9]+',s)).lower();audit={}
for f in sorted((C/'docs/labs/pr12/dk/reports').glob('*.docx')):
 d=Document(f);pdf=E/'render'/f.stem/(f.stem+'.pdf');pd=pdfplumber.open(pdf);pages=[p.extract_text() or '' for p in pd.pages];flat=norm('\n'.join(pages));missing=[];checked=0
 for typ,text in [('paragraph',p.text) for p in d.paragraphs]+[('table_cell',c.text) for t in d.tables for row in t.rows for c in row.cells]:
  s=norm(text)
  if len(s)<8:continue
  checked+=1
  if s not in flat:missing.append({'kind':typ,'text':text})
 with zipfile.ZipFile(f) as z:
  assert z.testzip() is None
  for name in z.namelist():
   if name.endswith(('.xml','.rels')):etree.fromstring(z.read(name))
 audit[f.name]={'sha256':sha(f),'zip_xml':'PASS','tables':len(d.tables),'table_rows':[len(t.rows) for t in d.tables],'embedded_media':len(d.inline_shapes),'pdf_sha256':sha(pdf),'pdf_pages':len(pages),'all_pages_nonempty':all(len(t.strip())>10 for t in pages),'checked_text_units':checked,'text_units_not_exact_contiguous_in_pdf':missing,'text_check_note':'Normalized paragraph/cell strings compared with extracted PDF text; table cells are interleaved across columns in PDF and can legitimately fail a contiguous-string test. XML/table counts and full text remain available.','page_text':[{'page':n+1,'text_chars':len(t),'start':t[:180],'png_sha256':sha(pdf.parent/f'page-{n+1}.png')} for n,t in enumerate(pages)],'visual_review':'ALL_PAGES_INSPECTED_IN_CONTACT_SHEETS; affected pages also inspected separately','layout_findings':['DK-REV-002'],'toc_note':'0302 reports have an uncached TOC field and only explanatory text on page 2; explicitly deferred by author, remains a finalization action.' if f.name.startswith('0302') else 'No TOC section in these exports'}
 pd.close()
(E/'document-review.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n');print({k:{'pages':v['pdf_pages'],'unmatched_text_units':len(v['text_units_not_exact_contiguous_in_pdf'])} for k,v in audit.items()})
